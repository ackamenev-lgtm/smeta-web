"""
Умное чтение смет: приоритет — шапка с итогами, потом позиции.
"""
import io
import re


def read_smeta_file(file_bytes: bytes, ext: str) -> str:
    if ext == '.docx':
        text = _read_docx(file_bytes)
    else:
        text = _read_xlsx(file_bytes)
    return _smart_compress(text)


def _read_docx(file_bytes: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(file_bytes))
    parts = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            unique = []
            for c in cells:
                if not unique or c != unique[-1]:
                    unique.append(c)
            line = ' | '.join(filter(None, unique))
            if line:
                parts.append(line)
    return '\n'.join(parts)


def _read_xlsx(file_bytes: bytes) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    parts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        parts.append(f'=== Лист: {sheet_name} ===')
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() if c is not None else '' for c in row]
            line = ' | '.join(filter(None, cells))
            if line:
                parts.append(line)
    return '\n'.join(parts)


def _smart_compress(raw: str) -> str:
    """
    Трёхуровневая приоритизация строк сметы.
    Числа в формате "1 797 442" (с пробелами) тоже детектируются.
    """
    lines = raw.split('\n')
    P1, P2, P3 = [], [], []
    seen = set()

    # Числа: как обычные (4+ цифры подряд), так и с пробелами-разделителями
    RE_NUM = re.compile(r'\d{4,}|\d{1,3}(?:\s\d{3})+(?:[,\.]\d+)?')

    RE_SEC_TOTAL = re.compile(
        r'итого по разделу|итого прямые затраты по разделу|'
        r'итого б.?ндс|итого без ндс|всего по смете', re.I
    )
    RE_STRUCT = re.compile(
        r'материальные ресурсы|оплата труда рабочих|оплата труда машинистов|'
        r'средства на оплату труда|эксплуатация машин|перевозка|перевоз|'
        r'накладные расходы|сметная прибыль|сметная стоимость|'
        r'строительных работ|монтажных работ|прочих затрат|'
        r'ндс\s*\d|ндс\s*22|нмц(?!\w)|нмцк', re.I
    )
    RE_SECTION = re.compile(r'^раздел\s+\d+\.', re.I)
    RE_CODE    = re.compile(r'ГЭСН|ГЭСНр|ГЭСНм|ФСН|\d{2}-\d{2}-\d{3}', re.I)
    RE_MAT     = re.compile(r'^\d{2}\.\d|\bпрайс\b', re.I)

    for line in lines:
        t = line.strip()
        if not t or len(t) < 4:
            continue
        key = t[:200]
        if key in seen:
            continue
        seen.add(key)
        t_cut = t[:320]
        has_num = bool(RE_NUM.search(t))

        if RE_SEC_TOTAL.search(t) or RE_STRUCT.search(t) or RE_SECTION.match(t):
            P1.append(t_cut)
        elif RE_CODE.search(t) and has_num:
            P2.append(t_cut)
        elif (RE_MAT.search(t) or 'прайс' in t.lower()) and has_num:
            P3.append(t_cut)
        elif has_num and len(t) < 220:
            P2.append(t_cut)

    result_lines = P1[:]
    remaining = 42000 - sum(len(l) for l in result_lines)
    for l in P2:
        if remaining <= 0:
            break
        result_lines.append(l)
        remaining -= len(l) + 1
    for l in P3:
        if remaining <= 0:
            break
        result_lines.append(l)
        remaining -= len(l) + 1

    result = '\n'.join(result_lines)
    return result[:44000] if len(result) > 44000 else result
