"""
Генерация Excel-книги проекта через openpyxl.
10 листов с полным форматированием.
"""
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as GCL

# ── Цвета ────────────────────────────────────────────────────────────────────
C_HDR   = "1F4E79"
C_SEC   = "BDD7EE"
C_SUB   = "DDEBF7"
C_RED   = "C00000"
C_GREEN = "375623"
C_CHEAP = "E2EFDA"
C_PRIC  = "FCE4D6"
C_EQUAL = "FFF2CC"

_thin = Side(style='thin', color='AAAAAA')
BRD   = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def N(v) -> int:
    try: return int(round(float(v or 0)))
    except: return 0

def F(v, dec=2):
    try: return round(float(v or 0), dec)
    except: return 0

def pct(a, b) -> str:
    try: return f"{round(a/b*100)}%" if b else "0%"
    except: return "0%"

def fill(color): return PatternFill("solid", start_color=color)
def font(bold=False, size=10, color="000000"):
    return Font(name="Arial", bold=bold, size=size, color=color)
def align(ha="left", v="center", wrap=False):
    return Alignment(horizontal=ha, vertical=v, wrap_text=wrap)


def _cell(ws, row, col, value=None, fmt=None, bold=False, size=10, color="000000",
          bg=None, ha="left", wrap=False, h=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = font(bold=bold, size=size, color=color)
    if bg: c.fill = fill(bg)
    c.alignment = align(ha, wrap=wrap)
    if fmt: c.number_format = fmt
    c.border = BRD
    if h: ws.row_dimensions[row].height = h
    return c


def _merge(ws, row, c1, c2, value=None, bold=False, size=10, color="000000",
           bg=None, ha="center", wrap=False, h=None):
    ws.merge_cells(f"{GCL(c1)}{row}:{GCL(c2)}{row}")
    c = _cell(ws, row, c1, value, bold=bold, size=size, color=color, bg=bg, ha=ha, wrap=wrap, h=h)
    for col in range(c1+1, c2+1):
        cs = ws.cell(row=row, column=col)
        if bg: cs.fill = fill(bg)
        cs.border = BRD
    return c


def _set_cols(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[GCL(i)].width = w


def _add_sheet(wb, name, tab_color):
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = tab_color
    return ws


def _banner(ws, row, cols, text, bold=True, size=10, bg=C_HDR, color="FFFFFF", h=26):
    _merge(ws, row, 1, cols, text, bold=bold, size=size, color=color, bg=bg, ha="center", h=h)


def _headers(ws, row, headers, bg=C_HDR, h=36):
    for ci, h_text in enumerate(headers, 1):
        c = ws.cell(row=row, column=ci, value=h_text)
        c.font = font(bold=True, color="FFFFFF")
        c.fill = fill(bg)
        c.alignment = align("center", wrap=True)
        c.border = BRD
    ws.row_dimensions[row].height = h


def _sec_row(ws, row, cols, text):
    _merge(ws, row, 1, cols, f"  {text}", bold=True, bg=C_SEC, ha="left", h=20)


def _tot_row(ws, row, cols, label, vals: dict, bg=C_HDR, color="FFFFFF"):
    _merge(ws, row, 1, 1, label, bold=True, color=color, bg=bg, ha="left", h=26)
    for col in range(2, cols+1):
        c = ws.cell(row=row, column=col)
        c.fill = fill(bg); c.border = BRD
        c.font = font(bold=True, color=color, size=11)
        if col in vals:
            c.value = vals[col]
            c.number_format = '#,##0'
            c.alignment = align("right")


# ════════════════════════════════════════════════════════════════════════
def generate_excel_book(data: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    proj  = data.get('project', {})
    st    = data.get('smeta_structure', {})
    secs  = data.get('sections', [])
    mkt   = data.get('market', {})
    msecs = mkt.get('sections', [])

    TOTAL = N(st.get('total_no_vat'))
    MAT   = N(st.get('materials'))
    OT    = N(st.get('ot_workers'))
    OTM   = N(st.get('otm_machines'))
    EM    = N(st.get('machine_op'))
    TR    = N(st.get('transport'))
    NR    = N(st.get('overhead'))
    SP    = N(st.get('profit'))
    VAT   = N(st.get('vat'))
    NMC   = N(st.get('nmc'))
    VAT_R = st.get('vat_rate', 0.22)

    MK_MAT = N(mkt.get('materials_no_vat'))
    MK_WRK = N(mkt.get('works_total'))
    SELF   = MK_MAT + MK_WRK + OTM + EM + TR
    ECO    = N(mkt.get('net_economy'))
    ECO_P  = round(float(mkt.get('net_economy_pct', 0)), 1)

    obj  = f"{proj.get('name','')} — {proj.get('address','')}"
    meta = f"ИКЗ: {proj.get('ikz','')}  |  {proj.get('defect_date','')}  |  {proj.get('price_period','')}"

    # ── 1. Сводный бюджет ────────────────────────────────────────────────
    ws1 = _add_sheet(wb, "📊 Сводный бюджет", C_RED)
    _set_cols(ws1, [38,18,18,18,10,34])
    R = 1
    _banner(ws1, R, 6, "СВОДНЫЙ БЮДЖЕТ ПРОЕКТА", size=14, h=36); R+=1
    _banner(ws1, R, 6, obj, bg="D6E4F0", color="000000", bold=True, size=10, h=20); R+=1
    _banner(ws1, R, 6, meta, bg="EEF4FA", color="444444", bold=False, size=8, h=18); R+=1
    _headers(ws1, R, ["Наименование","По смете б/НДС, руб.","При самост.орг. б/НДС, руб.","Разница, руб.","Разн. %","Примечание"]); ws1.freeze_panes=f"A{R+1}"; R+=1

    def drow(label, sm, mk, note="", bold=False, bg=None):
        nonlocal R
        diff = sm - mk if mk is not None else None
        p = pct(diff, sm) if diff is not None else ""
        if bg is None and diff is not None and sm:
            bg = C_CHEAP if diff>sm*0.03 else C_PRIC if diff<-sm*0.03 else C_EQUAL
        bg = bg or "FFFFFF"
        for ci,v in enumerate([label,sm,mk,diff,p,note],1):
            c = ws1.cell(row=R, column=ci, value=v)
            c.font=font(bold=bold); c.fill=fill(bg); c.border=BRD
            c.alignment=align("right" if ci in (2,3,4) else "left", wrap=(ci in (1,6)))
            if ci in (2,3,4) and v is not None: c.number_format='#,##0'
        ws1.row_dimensions[R].height=22; R+=1

    _sec_row(ws1, R, 6, "I. НМЦ КОНТРАКТА (по разделам)"); R+=1
    for sec in secs: drow(f"  {sec['name']}", N(sec.get('total')), None, bg="FFFFFF")
    _tot_row(ws1,R,6,"  ИТОГО без НДС",{2:TOTAL},bg="2F75B6"); R+=1
    drow(f"  НДС {round(VAT_R*100)}%", VAT, None, bg="D6E4F0")
    _tot_row(ws1,R,6,"НМЦ КОНТРАКТА (с НДС)",{2:NMC},bg=C_RED); R+=2

    _sec_row(ws1,R,6,"II. СТРУКТУРА ЗАТРАТ ПО СМЕТЕ"); R+=1
    for label,val in [("  Материалы",MAT),("  ОТ рабочих",OT),("  ОТм машинистов",OTM),
                      ("  Эксплуатация машин",EM),("  Перевозка",TR),
                      ("  Накладные расходы (НР)",NR),("  Сметная прибыль (СП)",SP)]:
        drow(label, val, None, bg="FFFFFF")
    _tot_row(ws1,R,6,"  ИТОГО",{2:TOTAL},bg="2F75B6"); R+=2

    _sec_row(ws1,R,6,"III. СМЕТА vs САМОСТОЯТЕЛЬНАЯ ОРГАНИЗАЦИЯ"); R+=1
    drow("  Материалы",MAT,MK_MAT,"petrovich/lemanapro ÷ 1.22")
    drow("  ОТ рабочих ↔ Рынок работ (бригады)",OT,MK_WRK,"Норматив ГЭСН vs расценки бригад",bold=True)
    drow("  Накладные расходы (НР)",NR,0,"При самост.орг. = 0",bold=True)
    drow("  Сметная прибыль (СП)",SP,0,"При самост.орг. = 0",bold=True)
    drow("  Прочее (ОТм, ЭМ, перевозка)",OTM+EM+TR,OTM+EM+TR,"Одинаково")
    _tot_row(ws1,R,6,"ИТОГО без НДС",{2:TOTAL,3:SELF,4:ECO,5:f"{ECO_P}%"},bg=C_GREEN); R+=1
    drow("  НДС",VAT,round(MK_MAT*0.22),bg="D6E4F0")
    _tot_row(ws1,R,6,"ИТОГО С НДС",{2:NMC,3:SELF+round(MK_MAT*0.22)},bg=C_RED); R+=2

    _sec_row(ws1,R,6,"IV. РАСШИФРОВКА ЭКОНОМИИ"); R+=1
    for label,val in [("+ Экономия на материалах",MAT-MK_MAT),
                      ("+ Отказ от НР подрядчика",NR),
                      ("+ Отказ от сметной прибыли",SP),
                      ("− Рынок работ > норм.ОТ",-(MK_WRK-OT))]:
        drow(label, val, None, bg=C_CHEAP if val>=0 else C_PRIC)
    _tot_row(ws1,R,6,"= ИТОГОВАЯ ЭКОНОМИЯ (без НДС)",{2:ECO,3:f"{ECO_P}%"},bg=C_GREEN)

    # ── 2. Оглавление ───────────────────────────────────────────────────
    ws2 = _add_sheet(wb, "🗂️ Оглавление", "FF0000")
    _set_cols(ws2, [6,34,50])
    _banner(ws2, 1, 3, f"КНИГА ПРОЕКТА — {proj.get('name','')}", h=32)
    _headers(ws2, 2, ["№","Лист","Содержание"], h=24)
    nav = [
        (1,"📊 Сводный бюджет","НМЦ, структура затрат, смета vs рынок, расшифровка экономии"),
        (2,"🗂️ Оглавление","Навигация по книге"),
        (3,"📋 Дефектная ведомость","Все работы с объёмами и стоимостью"),
        (4,"📋 Сводка по разделам","Итоги и доли разделов"),
        (5,"👷 ФОТ (оплата труда)","ОТ рабочих: ОТ на ед.изм. и итого"),
        (6,"🧱 Ведомость материалов","Все материалы: код, наименование, ед., кол-во, цена"),
        (7,"💰 Сравнение цен матер.","Каждый материал: смета vs рынок ÷ 1.22"),
        (8,"💰 Сводка цен матер.","Итоги по разделам: экономия на материалах"),
        (9,"⚒️ Стоимость работ","Каждая работа: ОТ смета vs расценка бригады"),
        (10,"⚒️ Итоговое сравнение","ОТ↔Рынок | НР+СП=0 | Расшифровка экономии"),
    ]
    for i,(num,sheet,desc) in enumerate(nav,3):
        bg = "FFFDE7" if num==1 else "EEF4FA" if i%2==0 else "FFFFFF"
        for ci,v in enumerate([num,sheet,desc],1):
            c=ws2.cell(row=i,column=ci,value=v)
            c.font=font(bold=(num==1)); c.fill=fill(bg); c.border=BRD
            c.alignment=align(wrap=(ci==3))
        ws2.row_dimensions[i].height=22

    # ── 3. Дефектная ведомость ───────────────────────────────────────────
    ws3 = _add_sheet(wb, "📋 Дефектная ведомость", C_HDR)
    _set_cols(ws3, [7,22,54,9,10,16])
    R3=1
    _banner(ws3,R3,6,"ДЕФЕКТНАЯ ВЕДОМОСТЬ"); R3+=1
    _banner(ws3,R3,6,obj,bg="D6E4F0",color="000000",bold=True,size=10,h=20); R3+=1
    _banner(ws3,R3,6,meta,bg="EEF4FA",color="444444",bold=False,size=8,h=18); R3+=1
    _headers(ws3,R3,["№ п/п","Шифр норматива","Наименование работ","Ед. изм.","Объём","Стоимость, руб."]); ws3.freeze_panes=f"A{R3+1}"; R3+=1

    for sec in secs:
        _merge(ws3,R3,1,6,f"  {sec['name']}",bold=True,bg=C_SEC,ha="left",h=20); R3+=1
        for pos in sec.get('positions',[]):
            for ci,v in enumerate([pos.get('num',''),pos.get('code',''),pos.get('name',''),
                                    pos.get('unit',''),F(pos.get('qty',0),3),N(pos.get('total',0))],1):
                c=ws3.cell(row=R3,column=ci,value=v)
                c.font=font(); c.border=BRD
                c.alignment=align("right" if ci in (5,6) else "center" if ci==1 else "left",wrap=(ci==3))
                if ci==6: c.number_format='#,##0'
            ws3.row_dimensions[R3].height=30; R3+=1
        sec_tot=N(sec.get('total'))
        _merge(ws3,R3,1,5,f"  Итого раздел: {sec['name']}",bold=True,bg=C_SUB,ha="right",h=20)
        c=ws3.cell(row=R3,column=6,value=sec_tot)
        c.fill=fill(C_SUB); c.border=BRD; c.number_format='#,##0'; c.font=font(bold=True); c.alignment=align("right"); R3+=1; R3+=1
    for lbl,val,bg in [("ИТОГО без НДС",TOTAL,C_HDR),(f"НДС {round(VAT_R*100)}%",VAT,"2F75B6"),("НМЦ (с НДС)",NMC,C_RED)]:
        _merge(ws3,R3,1,5,lbl,bold=True,color="FFFFFF",bg=bg,ha="right",h=24)
        c=ws3.cell(row=R3,column=6,value=val); c.fill=fill(bg); c.border=BRD
        c.number_format='#,##0'; c.font=font(bold=True,color="FFFFFF",size=11); c.alignment=align("right"); R3+=1

    # ── 4. Сводка по разделам ────────────────────────────────────────────
    ws4 = _add_sheet(wb,"📋 Сводка по разделам","2F75B6")
    _set_cols(ws4,[6,42,20,12])
    _banner(ws4,1,4,"СВОДКА ПО РАЗДЕЛАМ")
    _headers(ws4,2,["№","Раздел","Стоимость б/НДС, руб.","Доля, %"],h=28)
    for i,sec in enumerate(secs,3):
        tot=N(sec.get('total'))
        for ci,v in enumerate([i-2,sec['name'],tot,pct(tot,TOTAL)],1):
            c=ws4.cell(row=i,column=ci,value=v)
            c.font=font(); c.border=BRD
            c.alignment=align("center" if ci in (1,4) else "right" if ci==3 else "left")
            if ci==3: c.number_format='#,##0'
        ws4.row_dimensions[i].height=20
    for lbl,val,bg in [("Итого без НДС",TOTAL,C_SUB),(f"НДС {round(VAT_R*100)}%",VAT,C_SUB),("НМЦ (с НДС)",NMC,C_RED)]:
        r=3+len(secs)+([0,1,2][["Итого без НДС",f"НДС {round(VAT_R*100)}%","НМЦ (с НДС)"].index(lbl)])
        for ci,v in enumerate([None,lbl,val,None],1):
            c=ws4.cell(row=r,column=ci,value=v)
            c.fill=fill(bg); c.border=BRD; c.font=font(bold=True,color="FFFFFF" if bg==C_RED else "000000")
            c.alignment=align("right" if ci==3 else "left")
            if ci==3 and isinstance(v,int): c.number_format='#,##0'
        ws4.row_dimensions[r].height=24

    # ── 5. ФОТ ──────────────────────────────────────────────────────────
    ws5 = _add_sheet(wb,"👷 ФОТ (оплата труда)","375623")
    _set_cols(ws5,[7,20,52,8,10,14,16])
    _banner(ws5,1,7,"ДЕФЕКТНАЯ ВЕДОМОСТЬ — ЗАТРАТЫ НА ОПЛАТУ ТРУДА (ФОТ)")
    _banner(ws5,2,7,obj,bg="D6E4F0",color="000000",bold=True,size=10,h=20)
    _headers(ws5,3,["№","Шифр","Наименование работы","Ед.","Объём","ОТ на ед.изм., руб.","ОТ итого, руб."]); ws5.freeze_panes="A4"
    R5=4
    for sec in secs:
        _merge(ws5,R5,1,7,f"  {sec['name']}",bold=True,bg=C_SEC,ha="left",h=20); R5+=1
        for pos in sec.get('positions',[]):
            qty=F(pos.get('qty',0),3) or 1
            ot_t=N(pos.get('ot_total',0))
            ot_per=round(ot_t/qty,2) if qty else 0
            for ci,v in enumerate([pos.get('num',''),pos.get('code',''),pos.get('name',''),
                                    pos.get('unit',''),qty,ot_per,ot_t],1):
                c=ws5.cell(row=R5,column=ci,value=v)
                c.font=font(); c.border=BRD
                c.alignment=align("right" if ci in (5,6,7) else "center" if ci==1 else "left",wrap=(ci==3))
                if ci in (6,7): c.number_format='#,##0.00' if ci==6 else '#,##0'
            ws5.row_dimensions[R5].height=30; R5+=1
        sec_ot=sum(N(p.get('ot_total',0)) for p in sec.get('positions',[]))
        _merge(ws5,R5,1,6,f"  Итого ОТ: {sec['name']}",bold=True,bg=C_SUB,ha="right",h=20)
        c=ws5.cell(row=R5,column=7,value=sec_ot); c.fill=fill(C_SUB); c.border=BRD
        c.number_format='#,##0'; c.font=font(bold=True); c.alignment=align("right"); R5+=1; R5+=1
    for lbl,val in [("ИТОГО ОТ рабочих",OT),("ОТм машинистов (справочно)",OTM),("ИТОГО ФОТ",OT+OTM)]:
        _merge(ws5,R5,1,6,lbl,bold=True,color="FFFFFF",bg=C_HDR,ha="right",h=24)
        c=ws5.cell(row=R5,column=7,value=val); c.fill=fill(C_HDR); c.border=BRD
        c.number_format='#,##0'; c.font=font(bold=True,color="FFFFFF",size=11); c.alignment=align("right"); R5+=1

    # ── 6. Ведомость материалов ──────────────────────────────────────────
    ws6 = _add_sheet(wb,"🧱 Ведомость материалов","833C11")
    _set_cols(ws6,[20,7,18,46,9,10,14,16])
    _banner(ws6,1,8,"ВЕДОМОСТЬ МАТЕРИАЛОВ")
    _banner(ws6,2,8,obj,bg="D6E4F0",color="000000",bold=True,size=10,h=20)
    _headers(ws6,3,["Раздел","Поз.","Код","Наименование материала","Ед.","Кол-во","Цена б/НДС","Стоимость"]); ws6.freeze_panes="A4"
    R6=4
    for sec in secs:
        sec_mat_tot=0
        for mat in sec.get('materials',[]):
            tot=N(mat.get('total',0)); sec_mat_tot+=tot
            for ci,v in enumerate([sec['name'],'',mat.get('code',''),mat.get('name',''),
                                    mat.get('unit',''),F(mat.get('qty',0),3),F(mat.get('price',0),2),tot],1):
                c=ws6.cell(row=R6,column=ci,value=v)
                c.font=font(); c.border=BRD
                c.alignment=align("right" if ci in (6,7,8) else "left",wrap=(ci==4))
                if ci in (7,8): c.number_format='#,##0.00' if ci==7 else '#,##0'
            ws6.row_dimensions[R6].height=28; R6+=1
        if sec_mat_tot:
            _merge(ws6,R6,1,7,f"  Итого материалы {sec['name']}",bold=True,bg=C_SUB,ha="right",h=20)
            c=ws6.cell(row=R6,column=8,value=sec_mat_tot); c.fill=fill(C_SUB); c.border=BRD
            c.number_format='#,##0'; c.font=font(bold=True); c.alignment=align("right"); R6+=1; R6+=1
    _merge(ws6,R6,1,7,"ИТОГО материалы",bold=True,color="FFFFFF",bg=C_HDR,ha="right",h=24)
    c=ws6.cell(row=R6,column=8,value=MAT); c.fill=fill(C_HDR); c.border=BRD
    c.number_format='#,##0'; c.font=font(bold=True,color="FFFFFF",size=11); c.alignment=align("right")

    # ── 7. Сравнение цен материалов (построчно) ──────────────────────────
    ws7 = _add_sheet(wb,"💰 Сравнение цен матер.","C55A11")
    _set_cols(ws7,[6,36,8,10,14,14,16,14,14,14,9,20,30])
    _banner(ws7,1,13,"СРАВНЕНИЕ ЦЕН МАТЕРИАЛОВ — КАЖДАЯ ПОЗИЦИЯ")
    _merge(ws7,2,1,13,"Цены petrovich.ru / lemanapro.ru с НДС 22%, для сравнения ÷ 1.22  |  Май 2026",
           bg="EEF4FA",color="333333",bold=False,size=8,h=20)
    _headers(ws7,3,["№","Наименование материала","Ед.","Кол-во",
                    "Смет.цена б/НДС","Смет.стоим.",
                    "Рын.цена (с НДС 22%)","Рын.цена б/НДС (÷1.22)","Рын.стоим.",
                    "Экономия","%" ,"Источник/аналог","Раздел"]); ws7.freeze_panes="A4"
    R7=4; mat_num=0
    for sec in secs:
        for mat in sec.get('materials',[]):
            mat_num+=1
            sm_p=F(mat.get('price',0),2); sm_t=N(mat.get('total',0))
            mk_pv=N(mat.get('market_price_vat',0)); mk_p=F(mat.get('market_price',0),2)
            mk_t=N(mat.get('market_total',0)); eco=sm_t-mk_t
            bg=C_CHEAP if eco>sm_t*0.03 else C_PRIC if eco<-sm_t*0.03 else C_EQUAL
            if sm_t==0: bg="FFFFFF"
            for ci,v in enumerate([mat_num,mat.get('name',''),mat.get('unit',''),F(mat.get('qty',0),3),
                                    sm_p,sm_t,mk_pv,mk_p,mk_t,eco,pct(eco,sm_t) if sm_t else "—",
                                    mat.get('market_source',''),sec['name']],1):
                c=ws7.cell(row=R7,column=ci,value=v)
                c.font=font(); c.fill=fill(bg); c.border=BRD
                c.alignment=align("right" if ci in (4,5,6,7,8,9,10) else "left",wrap=(ci in (2,12)))
                if ci in (5,6,7,8,9,10): c.number_format='#,##0'
            ws7.row_dimensions[R7].height=30; R7+=1
    _merge(ws7,R7,1,5,"ИТОГО",bold=True,color="FFFFFF",bg=C_HDR,ha="right",h=24)
    for ci,v in [(6,MAT),(7,None),(8,None),(9,MK_MAT),(10,MAT-MK_MAT),(11,pct(MAT-MK_MAT,MAT))]:
        c=ws7.cell(row=R7,column=ci,value=v); c.fill=fill(C_HDR); c.border=BRD
        c.font=font(bold=True,color="FFFFFF",size=11); c.alignment=align("right")
        if ci in (6,9,10) and v is not None: c.number_format='#,##0'

    # ── 8. Сводка цен материалов ─────────────────────────────────────────
    ws8 = _add_sheet(wb,"💰 Сводка цен матер.","ED7D31")
    _set_cols(ws8,[36,18,18,18,10,30])
    _banner(ws8,1,6,"СВОДКА СРАВНЕНИЯ ЦЕН МАТЕРИАЛОВ ПО РАЗДЕЛАМ")
    _headers(ws8,2,["Раздел","Смета б/НДС, руб.","Рынок б/НДС, руб.","Экономия, руб.","Разн. %","Вывод"],h=32)
    for i,sec in enumerate(secs,3):
        ms=next((x for x in msecs if x.get('name')==sec['name']),{})
        sm=N(ms.get('mat_smeta') or sum(m.get('total',0) for m in sec.get('materials',[])))
        mk=N(ms.get('mat_market_no_vat') or sum(m.get('market_total',0) for m in sec.get('materials',[])))
        eco=sm-mk
        bg=C_CHEAP if eco>sm*0.03 else C_PRIC if eco<-sm*0.03 else C_EQUAL
        for ci,v in enumerate([sec['name'],sm,mk,eco,pct(eco,sm),"Закупка выгоднее" if eco>0 else "Дороже сметы"],1):
            c=ws8.cell(row=i,column=ci,value=v); c.font=font(); c.fill=fill(bg); c.border=BRD
            c.alignment=align("right" if ci in (2,3,4,5) else "left")
            if ci in (2,3,4): c.number_format='#,##0'
        ws8.row_dimensions[i].height=22
    tr8=3+len(secs)
    for ci,v in enumerate(["ИТОГО",MAT,MK_MAT,MAT-MK_MAT,pct(MAT-MK_MAT,MAT),f"Экономия {pct(MAT-MK_MAT,MAT)}"],1):
        c=ws8.cell(row=tr8,column=ci,value=v); c.fill=fill(C_HDR); c.border=BRD
        c.font=font(bold=True,color="FFFFFF",size=11); c.alignment=align("right" if ci>1 else "left")
        if ci in (2,3,4) and isinstance(v,int): c.number_format='#,##0'
    ws8.row_dimensions[tr8].height=26

    # ── 9. Стоимость работ (построчно) ────────────────────────────────────
    ws9 = _add_sheet(wb,"⚒️ Стоимость работ","7030A0")
    _set_cols(ws9,[6,46,8,10,14,14,14,14,14,9,24])
    _banner(ws9,1,11,"СРАВНЕНИЕ РАСЦЕНОК НА РАБОТЫ — КАЖДАЯ ПОЗИЦИЯ")
    _merge(ws9,2,1,11,"Смета ОТ — нормативная зарплата ГЭСН. Рыночная — только работа без материалов (otdelka-spb.ru, demontazh-centr.ru)",
           bg="EEF4FA",color="333333",bold=False,size=8,h=22)
    _headers(ws9,3,["№","Наименование работы","Ед.","Объём",
                    "ОТ на ед.изм., руб.","Смета ОТ итого",
                    "Рын.расценка руб./ед.","Рын.стоим.","Разница","Разн. %","Источник расценки"]); ws9.freeze_panes="A4"
    R9=4; pos_num=0
    for sec in secs:
        _merge(ws9,R9,1,11,f"  {sec['name']}",bold=True,bg=C_SEC,ha="left",h=20); R9+=1
        for pos in sec.get('positions',[]):
            pos_num+=1
            qty=F(pos.get('qty',0),3) or 1
            ot_t=N(pos.get('ot_total',0))
            ot_per=round(ot_t/qty,2) if qty else 0
            mr=N(pos.get('market_rate',0))
            mt=N(pos.get('market_total',0))
            diff=ot_t-mt
            bg=C_CHEAP if diff>ot_t*0.05 else C_PRIC if diff<-ot_t*0.05 else C_EQUAL
            if mr==0: bg="FFFFFF"
            for ci,v in enumerate([pos_num,pos.get('name',''),pos.get('unit',''),qty,
                                    ot_per,ot_t,mr,mt,diff,pct(diff,ot_t) if ot_t else "—",
                                    pos.get('work_source','otdelka-spb.ru')],1):
                c=ws9.cell(row=R9,column=ci,value=v); c.font=font(); c.fill=fill(bg); c.border=BRD
                c.alignment=align("right" if ci in (4,5,6,7,8,9) else "left",wrap=(ci==2))
                if ci in (5,6,7,8,9): c.number_format='#,##0.00' if ci==5 else '#,##0'
            ws9.row_dimensions[R9].height=32; R9+=1
    _merge(ws9,R9,1,5,"ИТОГО ОТ смета vs рыночные работы",bold=True,color="FFFFFF",bg=C_HDR,ha="right",h=28)
    for ci,v in [(6,OT),(7,None),(8,MK_WRK),(9,OT-MK_WRK),(10,pct(OT-MK_WRK,OT)),(11,None)]:
        c=ws9.cell(row=R9,column=ci,value=v); c.fill=fill(C_HDR); c.border=BRD
        c.font=font(bold=True,color="FFFFFF",size=11); c.alignment=align("right")
        if ci in (6,8,9) and v is not None: c.number_format='#,##0'

    # ── 10. Итоговое сравнение ────────────────────────────────────────────
    ws10 = _add_sheet(wb,"⚒️ Итоговое сравнение","4472C4")
    _set_cols(ws10,[44,18,18,16,9,36])
    _banner(ws10,1,6,"ИТОГОВОЕ СРАВНЕНИЕ: СМЕТА (ПОДРЯДЧИК) vs САМОСТОЯТЕЛЬНАЯ ОРГАНИЗАЦИЯ",size=12)
    _banner(ws10,2,6,"Материалы↔Материалы | ОТ рабочих↔Рынок работ | НР+СП при самост.орг.=0",
            bg="EEF4FA",color="333333",bold=False,size=8,h=22)
    _headers(ws10,3,["Статья затрат","По смете б/НДС, руб.","При самост.орг. б/НДС, руб.","Разница","Разн. %","Пояснение"]); ws10.freeze_panes="A4"
    R10=4

    def row10(lbl,sm,mk,note="",bold=False,bg=None):
        nonlocal R10
        diff=sm-mk if mk is not None else None
        p=pct(diff,sm) if diff is not None else ""
        if bg is None and diff is not None and sm:
            bg=C_CHEAP if diff>abs(sm)*0.03 else C_PRIC if diff<-abs(sm)*0.03 else C_EQUAL
        bg=bg or C_CHEAP
        for ci,v in enumerate([lbl,sm,mk,diff,p,note],1):
            c=ws10.cell(row=R10,column=ci,value=v)
            c.font=font(bold=bold); c.fill=fill(bg); c.border=BRD
            c.alignment=align("right" if ci in (2,3,4) else "left",wrap=(ci in (1,6)))
            if ci in (2,3,4) and v is not None: c.number_format='#,##0'
        ws10.row_dimensions[R10].height=24; R10+=1

    def tot10(lbl,sm,mk,note="",bg=C_HDR):
        nonlocal R10
        diff=sm-mk if mk else None; p=pct(diff,sm) if diff else ""
        for ci,v in enumerate([lbl,sm,mk,diff,p,note],1):
            c=ws10.cell(row=R10,column=ci,value=v)
            c.font=font(bold=True,color="FFFFFF",size=11); c.fill=fill(bg); c.border=BRD
            c.alignment=align("right" if ci>1 else "left",wrap=(ci==6))
            if ci in (2,3,4) and v is not None: c.number_format='#,##0'
            if ci==4: c.font=font(bold=True,color="FFFF00",size=12)
        ws10.row_dimensions[R10].height=28; R10+=1

    _sec_row(ws10,R10,6,"I. ПРЯМЫЕ ЗАТРАТЫ"); R10+=1
    row10("  Материалы",MAT,MK_MAT,"petrovich/lemanapro ÷ 1.22 (НДС 22%)")
    row10("  ОТ рабочих ↔ Рынок работ (бригады)",OT,MK_WRK,"Норматив ГЭСН vs расценки бригад",bold=True)
    row10("  ОТм машинистов",OTM,OTM,"Одинаково")
    row10("  Эксплуатация машин",EM,EM,"Одинаково")
    row10("  Перевозка (вывоз мусора)",TR,TR,"Одинаково")
    tot10("Итого прямые затраты",MAT+OT+OTM+EM+TR,MK_MAT+MK_WRK+OTM+EM+TR,"","2F75B6"); R10+=1

    _sec_row(ws10,R10,6,"II. КОСВЕННЫЕ ЗАТРАТЫ (только у подрядчика)"); R10+=1
    row10("  Накладные расходы (НР)",NR,0,"Управление/ИТР — при самост.орг.=0",bold=True)
    row10("  Сметная прибыль (СП)",SP,0,"Прибыль подрядчика — при самост.орг.=0",bold=True)
    tot10(f"Итого косвенные (НР+СП)",NR+SP,0,f"Полностью уходит подрядчику: {(NR+SP):,} руб.",C_GREEN); R10+=1

    _sec_row(ws10,R10,6,"III. ИТОГОВЫЙ БЮДЖЕТ"); R10+=1
    row10("Итого без НДС",TOTAL,SELF,f"Экономия {ECO_P}%",bold=True)
    row10("  НДС",VAT,round(MK_MAT*0.22))
    tot10("ИТОГО С НДС (НМЦ vs рыночный бюджет)",NMC,SELF+round(MK_MAT*0.22),"НМЦ контракта vs реальный бюджет",C_RED); R10+=1

    _sec_row(ws10,R10,6,"IV. РАСШИФРОВКА ЭКОНОМИИ"); R10+=1
    for lbl,val,note in [("+ Экономия на материалах",MAT-MK_MAT,"petrovich/lemanapro ÷ 1.22"),
                         ("+ Отказ от НР подрядчика",NR,"Управление/ИТР = 0"),
                         ("+ Отказ от сметной прибыли",SP,"Прибыль подрядчика = 0"),
                         ("− Рынок работ > норм.ОТ",-(MK_WRK-OT),"Бригады включают свои НР+прибыль")]:
        bg=C_CHEAP if val>=0 else C_PRIC
        for ci,v in enumerate([lbl,val,None,None,pct(val,TOTAL) if TOTAL else "—",note],1):
            c=ws10.cell(row=R10,column=ci,value=v); c.font=font(bold=True); c.fill=fill(bg); c.border=BRD
            c.alignment=align("right" if ci in (2,5) else "left",wrap=(ci in (1,6)))
            if ci==2 and v is not None: c.number_format='#,##0'
        ws10.row_dimensions[R10].height=24; R10+=1
    for ci,v in enumerate([f"= ИТОГОВАЯ ЭКОНОМИЯ (без НДС)",ECO,None,None,f"{ECO_P}%",""],1):
        c=ws10.cell(row=R10,column=ci,value=v)
        c.fill=fill(C_GREEN); c.border=BRD
        c.font=font(bold=True,color="FFFFFF",size=12 if ci in (2,5) else 11)
        c.alignment=align("right" if ci in (2,5) else "left")
        if ci==2: c.number_format='#,##0'
    ws10.row_dimensions[R10].height=34

    buf=io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf.read()
