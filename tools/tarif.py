"""Быстрый поиск по справочникам НЕОСТРОЙ без чтения xlsx целиком.

    python tools/tarif.py search перегородка ГКЛ 2 слоя      # подбор артикула
    python tools/tarif.py search окраска стен --msk -n 15
    python tools/tarif.py art 8.1.22 8.2.38                  # полная карточка + комплект
    python tools/tarif.py kit 7.2.25                         # комплект материалов
    python tools/tarif.py cat кабель                         # категории: отходы, наценки
    python tools/tarif.py rash ELM                           # расходники, % к оплате бригаде
    python tools/tarif.py rate OTD                           # ставки часа бригады
    python tools/tarif.py types                              # наценка по типу тендера
    python tools/tarif.py params                             # доставка, резерв, стоимость денег
    python tools/tarif.py dump-csv out.csv                   # весь тарификатор в CSV (для grep)
"""
import argparse
import csv
import hashlib
import pickle
import re
import sys
import warnings

import openpyxl

from paths import CACHE, MATERIALS, STAFF, TARIF

warnings.filterwarnings("ignore")


def _sig(*files):
    h = hashlib.sha1()
    for f in files:
        h.update(f.read_bytes())
    return h.hexdigest()[:16]


def _num(v):
    return v if isinstance(v, (int, float)) else None


def _build():
    wb = openpyxl.load_workbook(TARIF, data_only=True, read_only=True)
    works, sec = {}, ""
    for r in wb["Расценки"].iter_rows(min_row=5, values_only=True):
        art = r[0]
        if art is None:
            continue
        art = str(art).strip()
        if r[2] is None and r[1] is None:
            sec = art
            continue
        works[art] = dict(
            art=art, name=r[1], unit=r[2], new=bool(r[3]), spec=r[4], sec=sec,
            mkt_spb=_num(r[5]), mkt_msk=_num(r[6]), kp_spb=_num(r[7]), kp_msk=_num(r[8]),
            over_spb=_num(r[9]) or 0, over_msk=_num(r[10]) or 0,
            brig_spb=_num(r[11]), brig_msk=_num(r[12]), src=r[20], note=r[21],
        )
    for r in wb["Нормы и нормо-часы"].iter_rows(min_row=6, values_only=True):
        art = str(r[0]).strip() if r[0] is not None else ""
        if art in works:
            works[art].update(norm=_num(r[5]), norm_src=r[6], better_spb=r[16], better_msk=r[17])

    wm = openpyxl.load_workbook(MATERIALS, data_only=True, read_only=True)
    kits = {}
    cur = None
    for r in wm["Комплекты материалов"].iter_rows(min_row=4, values_only=True):
        if not any(r):
            continue
        if r[0]:
            cur = str(r[0]).strip()
        if cur and r[3]:
            kits.setdefault(cur, []).append(dict(
                material=r[3], cat=r[4], unit=r[5], rate=_num(r[6]), note=r[7]))
    cats = []
    for r in wm["Категории и наценки"].iter_rows(min_row=4, values_only=True):
        if r[0] and _num(r[1]) is not None:
            cats.append(dict(cat=r[0], waste=r[1], mk_sep=r[2], mk_turnkey=r[3], where=r[4], examples=r[5], note=r[6]))
    rash = {}
    for r in wm["Расходники к работам"].iter_rows(min_row=4, values_only=True):
        if r[0] and _num(r[2]) is not None:
            rash[r[0]] = dict(code=r[0], spec=r[1], pct=r[2], what=r[3], not_rash=r[4])
    types = []
    for r in wm["Тип тендера"].iter_rows(min_row=4, values_only=True):
        if r[0] and _num(r[1]) is not None:
            types.append(dict(type=r[0], mk_from=r[1], mk_to=r[2], margin_from=r[3], margin_to=r[4], ex=r[5]))
    params = []
    for r in wm["Параметры"].iter_rows(min_row=3, values_only=True):
        if r[0] and r[1] is not None:
            params.append(dict(name=r[0], value=r[1], note=r[2]))

    ws = openpyxl.load_workbook(STAFF, data_only=True, read_only=True)["Почасовка — ставки и план"]
    rates = {}
    for r in ws.iter_rows(min_row=4, values_only=True):
        if r[0] and _num(r[3]) is not None:
            rates[r[0]] = dict(code=r[0], spec=r[1], what=r[2], h_spb=r[3], h_msk=r[4],
                               mkt_h_spb=_num(r[9]), mkt_h_msk=_num(r[10]),
                               kp_h_spb=_num(r[11]), kp_h_msk=_num(r[12]))
    return dict(works=works, kits=kits, cats=cats, rash=rash, types=types, params=params, rates=rates)


def load():
    sig = _sig(TARIF, MATERIALS, STAFF)
    f = CACHE / f"spr_{sig}.pkl"
    if f.exists():
        return pickle.loads(f.read_bytes())
    data = _build()
    CACHE.mkdir(parents=True, exist_ok=True)
    f.write_bytes(pickle.dumps(data))
    return data


def _norm(s):
    s = str(s or "").lower().replace("ё", "е")
    return re.findall(r"[a-zа-я]+|\d+(?:[.,]\d+)?", s)


def _stem(t):
    return t[: max(4, len(t) - 2)] if len(t) > 5 and not t[0].isdigit() else t


SYN = {"гкл": "гипсокартон", "гипсокартон": "гкл", "гвл": "гипсоволокн", "кг": "керамогранит",
       "пвх": "поливинилхлорид", "вд": "водно", "ак": "акрил"}
DEMO = ("демонт", "разбор", "прорез", "снятие", "срезк", "вырезк")


def search(data, query, n=10):
    q = [_stem(t) for t in _norm(" ".join(query))]
    q_demo = any(s.startswith(DEMO) for s in q)
    res = []
    for w in data["works"].values():
        toks = _norm(w["name"])
        hit = 0
        for s in q:
            alt = SYN.get(s)
            if any(t.startswith(s) or (alt and t.startswith(alt)) for t in toks):
                hit += 1
        if hit:
            score = hit / len(q)
            if not q_demo and toks and toks[0].startswith(DEMO):
                score -= 0.3
            res.append((score, -len(toks) / 100, w))
    res.sort(key=lambda x: (x[0], x[1]), reverse=True)
    return [w for _, _, w in res[:n]]


def fmt(v):
    if v is None:
        return "—"
    if isinstance(v, float):
        return f"{v:.3f}".rstrip("0").rstrip(".") if abs(v) < 10 else f"{v:,.0f}".replace(",", " ")
    return str(v)


def line(w, msk=False):
    reg = "msk" if msk else "spb"
    over = w[f"over_{reg}"]
    flag = f" ПОДТЯНУТА +{over:.0%}" if over and over > 0.005 else ""
    return (f"{w['art']} | {w['name']} | {w['unit']} | {w['spec']} | КП {fmt(w[f'kp_{reg}'])} | "
            f"бригада {fmt(w[f'brig_{reg}'])} | рынок {fmt(w[f'mkt_{reg}'])} | норма {fmt(w.get('norm'))} ч"
            f"{flag}{' | комплект' if w['art'] in DATA['kits'] else ''}")


def main():
    global DATA
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("cmd")
    p.add_argument("args", nargs="*")
    p.add_argument("-n", type=int, default=10)
    p.add_argument("--msk", action="store_true", help="цены Москвы (по умолчанию СПб)")
    a = p.parse_args()
    DATA = load()
    reg = "Мск" if a.msk else "СПб"

    if a.cmd == "search":
        print(f"артикул | работа | ед. | код | цена КП {reg} | оплата бригаде | рынок | норма чел.-ч/ед.")
        for w in search(DATA, a.args, a.n):
            print(line(w, a.msk))
    elif a.cmd == "art":
        for art in a.args:
            w = DATA["works"].get(art)
            if not w:
                print(f"{art}: нет в тарификаторе")
                continue
            print(line(w, a.msk))
            print(f"   раздел: {w['sec']}; основание нормы: {w.get('norm_src')}; "
                  f"выгоднее отдать: {w.get('better_msk' if a.msk else 'better_spb')}; источник цены: {w['src']}")
            if w["note"]:
                print(f"   примечание: {w['note']}")
            for k in DATA["kits"].get(art, []):
                print(f"   комплект: {k['material']} | {k['cat']} | {fmt(k['rate'])} {k['unit']}/ед. | {k['note'] or ''}")
    elif a.cmd == "kit":
        for art in a.args:
            kit = DATA["kits"].get(art)
            print(f"{art}: " + ("нет типового комплекта — собрать по техкарте и записать в отчёт" if not kit else ""))
            for k in kit or []:
                print(f"   {k['material']} | {k['cat']} | {fmt(k['rate'])} {k['unit']}/ед. | {k['note'] or ''}")
    elif a.cmd == "cat":
        q = " ".join(a.args).lower()
        print("категория | отходы | наценка раздельная | под ключ | где проверяют")
        for c in DATA["cats"]:
            if not q or q in (c["cat"] + str(c["examples"])).lower():
                print(f"{c['cat']} | {c['waste']:.0%} | {c['mk_sep']:.0%} | {c['mk_turnkey']:.0%} | {c['where']}")
    elif a.cmd == "rash":
        for code, r in DATA["rash"].items():
            if not a.args or code in a.args:
                print(f"{code} | {r['spec']} | {r['pct']:.0%} к оплате бригаде | {r['what']}")
    elif a.cmd == "rate":
        print("код | специальность | ставка часа СПб | Мск | рынок нормо-часа СПб | Мск | нормо-час для заказчика СПб | Мск")
        for code, r in DATA["rates"].items():
            if not a.args or code in a.args:
                print(f"{code} | {r['spec']} | {fmt(r['h_spb'])} | {fmt(r['h_msk'])} | {fmt(r['mkt_h_spb'])} | "
                      f"{fmt(r['mkt_h_msk'])} | {fmt(r['kp_h_spb'])} | {fmt(r['kp_h_msk'])}")
    elif a.cmd == "types":
        for t in DATA["types"]:
            print(f"{t['type']} | наценка {fmt(t['mk_from'])}–{fmt(t['mk_to'])} | {t['ex'] or ''}")
    elif a.cmd == "params":
        for x in DATA["params"]:
            print(f"{x['name']} | {fmt(x['value'])} | {x['note'] or ''}")
    elif a.cmd == "dump-csv":
        out = a.args[0] if a.args else "tarif.csv"
        cols = ["art", "name", "unit", "spec", "kp_spb", "kp_msk", "brig_spb", "brig_msk",
                "mkt_spb", "mkt_msk", "over_spb", "over_msk", "norm", "sec"]
        with open(out, "w", newline="", encoding="utf-8") as f:
            wr = csv.writer(f)
            wr.writerow(cols)
            for w in DATA["works"].values():
                wr.writerow([w.get(c) for c in cols])
        print(out)
    else:
        print(__doc__)
        sys.exit(1)


if __name__ == "__main__":
    main()
